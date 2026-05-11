"""Личная переписка с Бишопом: /start, /что_ты_знаешь, готово, перенеси."""
from datetime import datetime
from zoneinfo import ZoneInfo

from aiogram import Bot, F, Router
from aiogram.filters import Command
from aiogram.types import (
    CallbackQuery,
    FSInputFile,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    KeyboardButton,
    Message,
    ReplyKeyboardMarkup,
    WebAppInfo,
)
from sqlalchemy import select
from sqlalchemy.orm import selectinload

from config import settings
from database import async_session_maker, Chat, Task, TaskAssignee, User
from services import claude_service, task_service, shop_tools
from handlers.messages import _ensure_user
from utils import log

router = Router()
TZ = ZoneInfo(settings.timezone)


PRICE_KEYWORDS = (
    "прайс", "посчит", "посчитай", "прикинь", "сколько будет", "сколько стоит",
    "цена ", "цены ", "позиц", "удали позиц", "покажи прайс", "обнови прайс",
    "добавь моносорт", "добавь микролот", "добавь бленд", "добавь blend",
    "новая позиция", "ещё позиция", "еще позиция",
    # КП по аренде кофейного оборудования
    "аренд", "арендую", "арендовать", "коммерческое предложение", "кп по аренде",
    "nimbus", "нимбус", "wpm", "zd-18", "zd 18", "кофемашин", "кофемолк",
)

GMAIL_KEYWORDS = (
    "почт", "gmail", "имейл", "имэйл", "email", "е-мейл", "e-mail",
    "инбокс", "inbox", "входящ",
    "дайджест", "сводк", "разбер", "разобрать", "разбор",
    "письм", "email", "писем",
    "что в почте", "почистить почту", "от тинькофф", "от сбер", "от альф",
    "от мтс", "от вб", "от ozon", "от озон", "от вайлдб",
    "что нового в", "что от ", "новые письма",
    "сколько спама", "сколько рекламы", "сколько уведомлений",
    "@gmail", "@mail", "@yahoo",
)

SHOP_KEYWORDS = (
    # Ozon Селлер — отчёты и аналитика по маркетплейсу
    "озон", "ozon", "озонe", "озоне", "озона", "озону",
    "маркетплейс", "маркетплейсе", "продажи на озоне", "продажи озон",
    "отчет по озон", "отчёт по озон", "отчет озон", "отчёт озон",
    "взаиморасчёт", "взаиморасчет", "сальдо",
    "комиссия озон", "штрафы озон", "лояльност",
    "позаказ", "реализаци", "выручка озон",
    "сколько в озоне", "сколько с озона",
    # Магазин (TMA / каталог)
    "магазин", "tma", "mini app", "miniapp",
    "добавь в магазин", "добавь в каталог", "в каталог магазина",
    "обнови фото", "поменяй фото", "загрузи фото", "это фото товара",
    "добавь фото", "сменить фото", "новое фото",
    "обнови описание", "поменяй описание",
    "опубликуй", "опубликуй магазин", "опубликовать в магазине",
    "товар в магазине", "позиция в магазине", "удали из магазина",
    "сток в магазине", "остаток в магазине",
    "убери из магазина", "убрать из магазина",
    # Ассортимент (общий прайс — молоко, сиропы, чай Althaus/Niktea, наборы)
    "сироп", "молоко", "топпинг", "ассортимент", "коэф",
    "althaus", "niktea", "альтхаус", "никти",
    "barline", "botanika", "ботаника", "herbarista", "гербариста",
    "sweetshot", "monin", "монин", "vedrenne", "ведренн",
    "набор tasteabrew", "ресторансия",
    # Файловые правки исходников (file-тулы, owner only)
    "открой файл", "покажи файл", "прочитай файл", "прочти файл",
    "поправь файл", "измени файл", "правка файла", "редактируй файл",
    "открой папку", "покажи папку", "содержимое папки",
    "обнови генератор", "поправь скрипт", "измени скрипт",
    "пересобери", "пересоберём", "пересоберем", "пересобрать",
    "убери из прайса", "убери из каталога", "убери из кп",
    "добавь в прайс", "добавь в каталог", "добавь в кп",
    "поправь прайс", "поправь каталог", "поправь кп",
    "поправь шрифт", "поправь цвет", "поменяй текст в",
    # Яндекс.Диск + PDF (yadisk_*, pdf_extract_pages)
    "яндекс", "yandex", "yadisk", "я.диск", "я диск",
    "диск", "облако", "на диске", "в облаке",
    "вытащи из pdf", "извлеки из pdf", "распакуй pdf",
    "pdf", "пдф", "карточки", "макет упаковки",
    "скачай файл", "забери файл", "возьми файл",
    # Остатки/сток магазина (sync_prices_stocks)
    "остатки", "остаток", "обнови остатки", "загрузи остатки",
    "новый сток", "сток обновился", "обнови сток", "залей остатки",
    "пересинкуй", "синкни остатки", "синк остатков",
    "прайс и остатки", "прайс остатки",
    # Roastberry Academy (students_tools, courses_tools)
    "академи", "academy", "школ",
    "курс 1", "курс 2", "курс 3", "курс 4", "курс 5",
    "курсы", "пройти курс", "доступ к курс",
    "пригласи", "пригласить", "приглашение",
    "ученик", "ученица", "ученики", "учеников",
    "аттестац", "аттестов", "сертификац",
    "прогрес", "прогресс ученик",
    "урок ", "уроки", "тест по ", "квиз",
    "magic-link", "magic link", "магик линк", "ссылку на курс",
    "открой курс", "открой доступ", "отзови доступ", "отозвать доступ",
    # Управление студентами / отправка курсов
    "отправь курс", "отправить курс", "выдай курс", "выдать курс",
    "пригласи на курс", "пригласить на курс", "приглашение на курс",
    "доступ на курс", "доступ к курс", "пусти на курс", "запиши на курс",
    "ученик", "ученики", "студент", "студенты",
    "академи", "academy",
    "продли доступ", "ограничь доступ", "временный доступ",
    "роастберри", "ростберри", "roastberry academy",
    "команда школы", "список учеников", "покажи учеников",
    "обнови урок", "поправь урок", "перепиши урок",
    "перерисуй картинку", "перерисуй иллюстрацию",
)


def _looks_like_price_request(text: str) -> bool:
    t = (text or "").lower()
    return any(kw in t for kw in PRICE_KEYWORDS)


def _looks_like_shop_request(text: str) -> bool:
    t = (text or "").lower()
    return any(kw in t for kw in SHOP_KEYWORDS)


def _looks_like_gmail_request(text: str) -> bool:
    t = (text or "").lower()
    return any(kw in t for kw in GMAIL_KEYWORDS)


SERVICE_KEYWORDS = (
    # Прямые упоминания сервиса
    "сервис", "service",
    "вызов", "вызовы", "вызвать мастер",
    "заявк", "запрос на ремонт", "ремонт",
    "сервисный бот", "service bot", "rbr_service_bot",
    # Поломки, проблемы с оборудованием
    "сломал", "не работа", "не варит", "не включа", "не запуск",
    "поломк", "поломал", "разобрался",
    "ошибка на дисплее", "ошибка экран", "код ошибки",
    "неисправн", "неполадк", "не функционирует",
    # Кофемашины — частая тема сервиса
    "кофемашина не", "кофемашины не", "паровик", "автомат не",
    "запасн", "запчаст", "замени деталь", "поменять деталь",
    # Запросы про мастера/техника
    "мастер прие", "мастер вые", "техник прие", "техник вые",
    "прислать мастер", "прислать техник",
    # Явные команды-вопросы про систему
    "как подать заявк", "как заявку", "куда заявк",
    "как сообщить о ремонт", "куда обратиться", "куда писать",
    "приложение сервис", "приложение мастер",
    # Контракты сервисной службы
    "франко", "алеф", "horeca", "хорека",
)


def _looks_like_service_request(text: str) -> bool:
    t = (text or "").lower()
    return any(kw in t for kw in SERVICE_KEYWORDS)


# Админ-команды владельца над сервисной службой (требуют tool-use в shop_chat,
# а НЕ выдачу плашки со ссылками). Имеют приоритет над SERVICE_KEYWORDS.
SERVICE_ADMIN_KEYWORDS = (
    # Перемещение/изменение существующих заявок
    "перенеси заявк", "перемести заявк", "переместить заявк",
    "поменяй контракт", "измени контракт",
    # Выдача кодов
    "выдай код", "сгенерируй код", "сделай приглашение", "code для техник",
    "код для техник", "код для менеджер", "пригласи техник", "пригласи менеджер",
    "сгенерируй приглашение", "выдай приглашение",
    # Статистика
    "покажи команду сервис", "статистика сервис", "статистика по сервис",
    "кто работает в сервис", "сколько кодов выдано", "пользователи сервис",
    "сводка сервис",
    # Email-routes
    "впредь от ", "впредь все письма", "сделай чтобы все от",
    "запомни отправител", "email route", "email-route",
    # Отправка ссылки
    "скинь ссылку на приложение", "отправь ссылку на приложение",
    "отправь @", "скинь @",
)


def _looks_like_service_admin_request(text: str) -> bool:
    t = (text or "").lower()
    return any(kw in t for kw in SERVICE_ADMIN_KEYWORDS)


def _service_links_keyboard():
    """Inline-кнопки со ссылками на сервисный бот @rbr_service_bot."""
    from aiogram.types import InlineKeyboardButton, InlineKeyboardMarkup
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(
            text="🚀 Открыть приложение",
            url="https://t.me/rbr_service_bot/service",
        )],
        [InlineKeyboardButton(
            text="🛠 Я техник",
            url="https://t.me/rbr_service_bot?start=technician",
        ),
         InlineKeyboardButton(
            text="📋 Я менеджер",
            url="https://t.me/rbr_service_bot?start=manager",
        )],
        [InlineKeyboardButton(
            text="☎️ Я клиент (подать заявку)",
            url="https://t.me/rbr_service_bot?start=client",
        )],
    ])


async def _send_long(message: Message, text: str):
    """Telegram режет на 4096; делим по 4000 на всякий."""
    if not text:
        await message.answer("(пустой ответ)")
        return
    chunk = 4000
    for i in range(0, len(text), chunk):
        await message.answer(text[i:i + chunk])


async def _send_shop_photos(message: Message, photos: list[dict]):
    """Отправляет фото или документы. PDF/XLSX уходят как document."""
    for ph in photos:
        path = ph["path"]
        is_doc = path.lower().endswith((".pdf", ".xlsx", ".docx", ".zip", ".csv"))
        try:
            if is_doc:
                await message.answer_document(
                    FSInputFile(path, filename=ph.get("filename")),
                    caption=ph.get("caption") or None,
                )
            else:
                await message.answer_photo(
                    FSInputFile(path),
                    caption=ph.get("caption") or None,
                )
        except Exception as e:
            log.error(f"send shop photo failed for {ph.get('path')}: {e}")
            await message.answer(f"Не смог отправить фото: {e}")


async def _send_price_files(message: Message, files: list[dict]):
    """Шлёт PDF/XLSX чистовики в личку."""
    for f in files:
        try:
            await message.answer_document(
                FSInputFile(f["path"], filename=f.get("filename")),
                caption=f.get("caption") or None,
            )
        except Exception as e:
            log.error(f"send_document failed for {f.get('path')}: {e}")
            await message.answer(f"Не смог отправить файл: {e}")


@router.message(F.chat.type == "private", Command(commands=["price", "prices"]))
async def cmd_price(message: Message):
    claude_service.reset_price_history(message.from_user.id)
    if message.from_user.id == settings.owner_telegram_id:
        await message.answer(
            "🧾 Прайс-режим (владелец). Что умею:\n"
            "• «посчитай моносорт Эфиопия Иргачеффе по $14»\n"
            "• «добавь микролот Кения Нямбени по $18»\n"
            "• «покажи прайс»\n"
            "• «удали позицию Бразилия Судан Руме анаэроб.»\n\n"
            "После расчёта спрошу подтверждение перед записью.\n"
            "Чтобы выйти — /price_off."
        )
    else:
        await message.answer(
            "🧾 Прайс. Спрашивайте:\n"
            "• «покажи прайс»\n"
            "• «есть ли в наличии Эфиопия Иргачеффе»\n"
            "• «сколько стоит Кения АА от 10 кг»\n\n"
            "Расчёт и добавление позиций — только у Дмитрия."
        )


@router.message(F.chat.type == "private", Command(commands=["price_off", "stop_price"]))
async def cmd_price_off(message: Message):
    claude_service.reset_price_history(message.from_user.id)
    await message.answer("Прайс-режим сброшен.")


def _all_dashboards_keyboard(user_id: int = 0) -> InlineKeyboardMarkup | None:
    """Клавиатура с кнопками дашбордов, доступных пользователю.

    Владелец видит все. Остальные — только те, на которые у них есть права
    (сейчас: prod_dashboard для prod_dashboard_user_id_set).
    """
    is_owner = (user_id == settings.owner_telegram_id)
    ops_url = (settings.ops_dashboard_url or "").strip() if is_owner else ""
    pl_url = (settings.pl_dashboard_url or "").strip() if is_owner else ""
    prod_url = (settings.prod_dashboard_url or "").strip() \
        if user_id in settings.prod_dashboard_user_id_set else ""
    rows = []
    if prod_url:
        rows.append([InlineKeyboardButton(
            text="🔥 Производство (обжарка 2026)",
            web_app=WebAppInfo(url=prod_url),
        )])
    if ops_url:
        rows.append([InlineKeyboardButton(
            text="📊 Операционный (1С + магазин + сервис)",
            web_app=WebAppInfo(url=ops_url),
        )])
    if pl_url:
        rows.append([InlineKeyboardButton(
            text="💰 Финансовый (P&L · 2021–2026)",
            web_app=WebAppInfo(url=pl_url),
        )])
    return InlineKeyboardMarkup(inline_keyboard=rows) if rows else None


@router.message(F.chat.type == "private",
                Command(commands=["dashboards", "дашборды", "панели", "dash", "панель"]))
async def cmd_all_dashboards(message: Message):
    """Единая точка входа во все дашборды. Для владельца — три, для прочих
    авторизованных — только те что разрешены."""
    uid = message.from_user.id
    kb = _all_dashboards_keyboard(uid)
    if not kb:
        # Молча игнорируем для не-владельцев без прав — чтобы не светить наличие.
        if uid != settings.owner_telegram_id:
            return
        await message.answer(
            "Дашборды пока не настроены. В .env Bishop задай "
            "<code>OPS_DASHBOARD_URL</code>, <code>PL_DASHBOARD_URL</code>, "
            "<code>PROD_DASHBOARD_URL</code>.",
            parse_mode="HTML",
        )
        return
    await message.answer(
        "Дашборды Roastberry\n"
        "Выбери, что открыть:",
        reply_markup=kb,
    )


@router.message(F.chat.type == "private", Command("start"))
async def cmd_start(message: Message):
    async with async_session_maker() as session:
        user = await _ensure_user(session, message.from_user)
        user.has_started_dm = True
        await session.commit()

    await message.answer(
        f"Привет, {message.from_user.first_name}! 👋\n\n"
        "Я Бишоп — помощник команды RBR.\n\n"
        "Теперь я смогу писать тебе напоминания по задачам.\n\n"
        "Что я умею:\n"
        "• Напоминаю о твоих задачах (заранее и в день дедлайна)\n"
        "• Принимаю твои ответы: \"готово\", \"перенеси на ...\"\n"
        "• Ищу по истории рабочих чатов\n\n"
        "Команды:\n"
        "/мои_задачи — список открытых задач\n"
        "/что_ты_знаешь — что я читаю и как работаю"
    )

    # Показываем дашборды сразу после /start всем у кого есть доступ.
    kb = _all_dashboards_keyboard(message.from_user.id)
    if kb:
        await message.answer("📂 Твои дашборды:", reply_markup=kb)


@router.message(F.chat.type == "private", Command(commands=["что_ты_знаешь", "help"]))
async def cmd_what_you_know(message: Message):
    async with async_session_maker() as session:
        result = await session.execute(select(Chat).where(Chat.is_active == True))
        chats = result.scalars().all()
        chat_list = "\n".join(f"• {c.title}" for c in chats) or "• (пока нет)"

    await message.answer(
        "🤖 Как я работаю\n\n"
        "Я читаю сообщения в рабочих чатах куда меня добавили:\n"
        f"{chat_list}\n\n"
        "❌ Я НЕ читаю:\n"
        "• Личные переписки сотрудников\n"
        "• Чаты куда меня не добавляли\n\n"
        "📋 Что я делаю:\n"
        "• Запоминаю задачи когда в чате пишут @bishoprb\n"
        "• Напоминаю исполнителям в личку\n"
        "• Ищу по истории чатов по запросу @bishoprb <вопрос>\n"
        "• Эскалирую просроченные задачи постановщику\n\n"
        "🔗 <b>Полезные команды:</b>\n"
        "• /мои_задачи — мои открытые задачи\n"
        "• /shop — режим магазина (поиск товаров, цены)\n"
        "• /service (или /сервис) — ссылки на бот сервисной службы\n"
        "• /digest — сводка по почте (если настроено)\n"
        "• /monitor — статус uptime-мониторов\n\n"
        "Я НЕ выдаю отчёты на конкретных людей. Я помощник, а не надзиратель.",
        parse_mode="HTML",
    )


@router.callback_query(F.data.startswith("prdone:"))
async def cb_personal_reminder_done(cb: CallbackQuery):
    """Кнопка «✅ Принял» под персональным напоминанием — закрываем задачу."""
    try:
        task_id = int(cb.data.split(":", 1)[1])
    except Exception:
        await cb.answer("Неверный формат", show_alert=False)
        return
    async with async_session_maker() as session:
        task = await task_service.get_task_by_id(session, task_id)
        if not task:
            await cb.answer("Задача не найдена", show_alert=False)
            return
        if task.creator_id != cb.from_user.id:
            await cb.answer("Это не твоё напоминание", show_alert=True)
            return
        await task_service.complete_task(session, task_id)
    # Убираем кнопку и помечаем что принял
    try:
        old_text = cb.message.html_text if cb.message else ""
        new_text = (old_text or "🗓 Напоминание") + "\n\n✅ <i>Принято</i>"
        await cb.message.edit_text(new_text[:4000], parse_mode="HTML", reply_markup=None)
    except Exception:
        pass
    await cb.answer("Принято — больше не напомню", show_alert=False)


@router.message(F.chat.type == "private",
                Command(commands=["ops", "опер", "операционный", "operations", "work"]))
async def cmd_ops_dashboard(message: Message):
    """Открывает операционный дашборд (1С + Магазин + Ozon + Сервис) в Telegram WebApp."""
    if message.from_user.id != settings.owner_telegram_id:
        return
    url = (settings.ops_dashboard_url or "").strip()
    if not url:
        await message.answer(
            "Операционный дашборд работает локально на 127.0.0.1:8083.\n\n"
            "Для открытия в Telegram нужен публичный HTTPS-URL. "
            "В .env Bishop задай:\n"
            "<code>ops_dashboard_url=https://&lt;твой-домен&gt;/ops_dashboard/</code>\n\n"
            "Сейчас можно открыть через SSH-туннель:\n"
            "<code>ssh -L 8083:127.0.0.1:8083 my-server</code>\n"
            "затем в браузере: http://127.0.0.1:8083/ops_dashboard/",
            parse_mode="HTML",
        )
        return
    kb = InlineKeyboardMarkup(inline_keyboard=[[
        InlineKeyboardButton(text="📊 Открыть операционный дашборд", web_app=WebAppInfo(url=url)),
    ]])
    await message.answer(
        "Операционный дашборд Roastberry\n"
        "Доступен только тебе. Внутри: продажи, маржа, остатки магазина, "
        "зерно, дебиторка, заказы магазина, Ozon, сервисная служба.",
        reply_markup=kb,
    )


@router.message(F.chat.type == "private",
                Command(commands=["prod", "production", "производство", "обжарка", "обжарки"]))
async def cmd_prod_dashboard(message: Message):
    """Открывает производственный дашборд (обжарка 2026) в Telegram WebApp.
    Доступ: владелец + PROD_DASHBOARD_USER_IDS (начальник производства и др.)."""
    if message.from_user.id not in settings.prod_dashboard_user_id_set:
        return
    url = (settings.prod_dashboard_url or "").strip()
    if not url:
        await message.answer(
            "Производственный дашборд работает локально на 127.0.0.1:8084.\n\n"
            "Для открытия в Telegram нужен публичный HTTPS-URL. В .env Bishop:\n"
            "<code>PROD_DASHBOARD_ENABLED=true</code>\n"
            "<code>PROD_DASHBOARD_URL=https://&lt;твой-домен&gt;/prod_dashboard/</code>\n\n"
            "Локально через SSH-туннель:\n"
            "<code>ssh -L 8084:127.0.0.1:8084 my-server</code>\n"
            "затем http://127.0.0.1:8084/prod_dashboard/",
            parse_mode="HTML",
        )
        return
    kb = InlineKeyboardMarkup(inline_keyboard=[[
        InlineKeyboardButton(text="🔥 Открыть производственный дашборд", web_app=WebAppInfo(url=url)),
    ]])
    await message.answer(
        "Производственный дашборд Roastberry · 2026\n"
        "Внутри: обжарки по ростерам, обжарщикам, графику, сверхурочка с разбором по дням и месяцам.",
        reply_markup=kb,
    )


@router.message(F.chat.type == "private", Command(commands=["pl", "финансы", "дашборд"]))
async def cmd_pl_dashboard(message: Message):
    """Открывает финансовый дашборд (P&L) в Telegram WebApp. Только для владельца."""
    if message.from_user.id != settings.owner_telegram_id:
        return
    url = (settings.pl_dashboard_url or "").strip()
    if not url:
        await message.answer(
            "Финансовый дашборд не настроен.\n\n"
            "В .env Bishop задайте:\n"
            "PL_DASHBOARD_ENABLED=true\n"
            "PL_DASHBOARD_URL=https://<твой-домен>/pl_dashboard/"
        )
        return
    kb = InlineKeyboardMarkup(inline_keyboard=[[
        InlineKeyboardButton(text="📊 Открыть финансовый дашборд", web_app=WebAppInfo(url=url)),
    ]])
    await message.answer(
        "Финансовый дашборд Roastberry P&L\n"
        "Доступен только тебе. Внутри: P&L, баланс, зарплаты, продажи, cash flow, склад, кредиты, оценка стоимости.",
        reply_markup=kb,
    )


# ── Естественные фразы про дашборды (без команд) ─────────────────────────
# Сработает раньше catch-all chat-handler. Только для владельца.
_DASH_KEYWORDS = (
    "дашборд", "дашбоард", "дашбоар", "дашбор", "дешборд",
    "dashboard",
)
_DASH_OPS_KEYWORDS = (
    "опер", "продаж", "склад", "выручк", "магазин", "ozon", "озон",
    "сервис", "заявк", "зерн", "дебитор", "остатк",
    "управленч", "операц",
)
_DASH_PL_KEYWORDS = (
    "финанс", "p&l", "пнл", "прибыл", "p&amp;l",
    "балан", "ebitda", "ебитда", "ебидта", "ебитд",
    "оценк", "valuation", "зп ", "зарпл", "кэшфлоу",
    "cashflow", "cash flow",
)
_DASH_PROD_KEYWORDS = (
    "производ", "обжарк", "обжарщ", "ростер", "сверхуроч",
    "production", "roast",
)


def _is_dashboard_intent(message: Message) -> bool:
    """Magic-filter: True только если сообщение похоже на просьбу открыть дашборд."""
    if not message.from_user or message.from_user.id != settings.owner_telegram_id:
        return False
    if message.chat.type != "private":
        return False
    text = (message.text or "").lower().strip()
    if not text or text.startswith("/"):
        return False
    return any(kw in text for kw in _DASH_KEYWORDS)


@router.message(_is_dashboard_intent)
async def msg_dashboard_intent(message: Message):
    """Свободные фразы про дашборд: «открой дашборд», «покажи операционный»,
    «открой p&l», «дашборд продаж» — шлёт кнопку Mini App.
    """
    text = (message.text or "").lower().strip()
    # Определяем какой дашборд
    is_pl = any(kw in text for kw in _DASH_PL_KEYWORDS)
    is_ops = any(kw in text for kw in _DASH_OPS_KEYWORDS)
    is_prod = any(kw in text for kw in _DASH_PROD_KEYWORDS)

    ops_url = (settings.ops_dashboard_url or "").strip()
    pl_url = (settings.pl_dashboard_url or "").strip()
    prod_url = (settings.prod_dashboard_url or "").strip()

    buttons = []
    if is_pl and pl_url:
        buttons.append([InlineKeyboardButton(
            text="💰 Финансовый (P&L)", web_app=WebAppInfo(url=pl_url),
        )])
    if is_ops and ops_url:
        buttons.append([InlineKeyboardButton(
            text="📊 Операционный", web_app=WebAppInfo(url=ops_url),
        )])
    if is_prod and prod_url:
        buttons.append([InlineKeyboardButton(
            text="🔥 Производственный", web_app=WebAppInfo(url=prod_url),
        )])
    # Если не уверен какой — все три
    if not buttons:
        if ops_url:
            buttons.append([InlineKeyboardButton(
                text="📊 Операционный", web_app=WebAppInfo(url=ops_url),
            )])
        if pl_url:
            buttons.append([InlineKeyboardButton(
                text="💰 Финансовый (P&L)", web_app=WebAppInfo(url=pl_url),
            )])
        if prod_url:
            buttons.append([InlineKeyboardButton(
                text="🔥 Производственный", web_app=WebAppInfo(url=prod_url),
            )])

    if not buttons:
        await message.answer(
            "Дашборды пока не настроены. Команды: <code>/ops</code>, <code>/финансы</code>, <code>/prod</code>.",
            parse_mode="HTML",
        )
        return

    kb = InlineKeyboardMarkup(inline_keyboard=buttons)
    picked = sum([is_pl, is_ops, is_prod])
    if picked == 1 and is_pl:
        title = "Финансовый дашборд (P&L)"
    elif picked == 1 and is_ops:
        title = "Операционный дашборд"
    elif picked == 1 and is_prod:
        title = "Производственный дашборд"
    else:
        title = "Выбери дашборд"
    await message.answer(title, reply_markup=kb)


@router.message(F.chat.type == "private", Command(commands=["id", "myid"]))
async def cmd_my_id(message: Message):
    """Показывает пользователю его telegram_id и username — нужно для приглашений в Академию."""
    user = message.from_user
    if not user:
        return
    parts = [
        f"🪪 <b>Твой Telegram-ID:</b> <code>{user.id}</code>",
    ]
    if user.username:
        parts.append(f"📛 <b>Username:</b> @{user.username}")
    if user.first_name:
        name = user.first_name + (f" {user.last_name}" if user.last_name else "")
        parts.append(f"👤 <b>Имя:</b> {name}")
    parts.append("")
    parts.append("Перешли это сообщение тому, кто хочет добавить тебя в Roastberry Academy.")
    await message.answer("\n".join(parts), parse_mode="HTML")


@router.message(F.chat.type == "private", Command("мои_задачи"))
async def cmd_my_tasks(message: Message):
    async with async_session_maker() as session:
        tasks = await task_service.get_pending_tasks_for_user(
            session, message.from_user.id
        )
    if not tasks:
        await message.answer("У тебя нет открытых задач. 🎉")
        return

    lines = ["📋 Твои открытые задачи:\n"]
    for t in tasks:
        lines.append(
            f"#{t.id} — {t.description}\n"
            f"⏰ До: {t.deadline.strftime('%d.%m %H:%M')}\n"
            f"👤 От: {t.creator.display_name}\n"
        )
    lines.append(
        "\nЧтобы закрыть: напиши \"готово #<номер>\" или просто опиши результат."
    )
    await message.answer("\n".join(lines))


@router.message(F.chat.type == "private", F.photo)
async def handle_private_photo(message: Message, bot: Bot):
    """Сотрудник прислал фото в личку — сохраняем в pending для shop_set_photo_from_telegram.
    Затем если в подписи есть ключ — сразу обрабатываем как магазин-запрос."""
    is_owner = message.from_user.id == settings.owner_telegram_id
    if not is_owner:
        await message.answer("Магазином управляет Дмитрий — фото от тебя я не сохраняю.")
        return
    try:
        # Берём самое большое фото
        file_id = message.photo[-1].file_id
        file = await bot.get_file(file_id)
        buf = await bot.download_file(file.file_path)
        image_bytes = buf.read() if hasattr(buf, "read") else bytes(buf)
        shop_tools.set_pending_photo(message.from_user.id, image_bytes)
        log.info(f"shop pending photo set for user {message.from_user.id}, {len(image_bytes)} bytes")
    except Exception as e:
        log.error(f"failed to download photo: {e}")
        await message.answer(f"Не смог скачать фото: {e}")
        return

    caption = (message.caption or "").strip()
    if not caption:
        await message.answer(
            "📸 Фото получено и ждёт. Теперь напиши к какому товару прикрепить, например:\n"
            "«это фото для Бразилия Серрадо 1 кг»\n"
            "или «обнови фото у NIKTEA Молочный Улун»."
        )
        return
    # Если есть подпись — сразу обрабатываем как shop-запрос
    try:
        answer, photos = await claude_service.shop_chat(
            caption + "\n\n[К сообщению приложено фото — используй shop_set_photo_from_telegram]",
            message.from_user.id, is_owner=True,
        )
    except Exception as e:
        log.error(f"shop_chat with photo failed: {e}")
        answer, photos = f"Ошибка: {e}", []
    if photos:
        await _send_shop_photos(message, photos)
    if answer:
        await _send_long(message, answer)


@router.message(F.chat.type == "private", F.document)
async def handle_private_document(message: Message, bot: Bot):
    """Принимаем документы. По типу маршрутизируем:
    - xlsx с 'остатки/прайс/ведомость' → синк магазина
    - pdf → pending для shop_set_photo_from_pending_pdf (страница как фото товара)
    - всё остальное — подсказка.
    """
    if message.from_user.id not in settings.shop_admin_id_set:
        await message.answer("Файлы для магазина принимаю только от админов.")
        return

    doc = message.document
    fname = (doc.file_name or "").strip()
    fname_low = fname.lower()

    # PDF → pending для shop_set_photo_from_pending_pdf
    if fname_low.endswith(".pdf") or (doc.mime_type or "").lower() == "application/pdf":
        try:
            file = await bot.get_file(doc.file_id)
            buf = await bot.download_file(file.file_path)
            pdf_bytes = buf.read() if hasattr(buf, "read") else bytes(buf)
            shop_tools.set_pending_pdf(message.from_user.id, pdf_bytes)
        except Exception as e:
            log.exception("pdf download failed")
            await message.answer(f"❌ Не смог скачать PDF: {e}")
            return
        caption = (message.caption or "").strip()
        if not caption:
            await message.answer(
                f"📄 PDF получен ({len(pdf_bytes)//1024} КБ). "
                f"Теперь напиши какой товар и какую страницу взять как фото, например:\n"
                f"«поставь первую страницу как фото у Колумбия Кастильо»\n"
                f"«страницу 3 на карточку Эфиопия Бомбе»",
                parse_mode="HTML",
            )
            return
        # Есть caption — сразу обрабатываем
        try:
            answer, photos = await claude_service.shop_chat(
                caption + "\n\n[К сообщению приложен PDF — используй shop_set_photo_from_pending_pdf]",
                message.from_user.id, is_owner=True,
            )
        except Exception as e:
            log.error(f"shop_chat with pdf failed: {e}")
            answer, photos = f"Ошибка: {e}", []
        if photos:
            await _send_shop_photos(message, photos)
        if answer:
            await _send_long(message, answer)
        return

    if not fname_low.endswith(".xlsx"):
        await message.answer(f"Принимаю только .xlsx и .pdf (а это {doc.mime_type or 'неизвестный формат'}).")
        return

    # Скачиваем в локальный путь TG-BOT под именем которое ждёт скрипт
    import shutil, subprocess
    from pathlib import Path

    looks_like_shop = (
        "остатк" in fname_low or "прайс" in fname_low or "ведомост" in fname_low
        or "stock" in fname_low or "price" in fname_low
    )
    if not looks_like_shop:
        await message.answer(
            f"📎 Получил <b>{fname}</b>, но не понял что это.\n"
            f"Жду xlsx где в имени есть «прайс», «остатки» или «ведомость».",
            parse_mode="HTML",
        )
        return

    # Скачиваем во временный файл, потом по содержимому решаем куда сохранить.
    try:
        file = await bot.get_file(doc.file_id)
        buf = await bot.download_file(file.file_path)
        data = buf.read() if hasattr(buf, "read") else bytes(buf)
    except Exception as e:
        log.exception("xlsx download failed")
        await message.answer(f"❌ Не смог скачать файл: {e}")
        return

    tmp = Path("/tmp/bishop_inbound.xlsx")
    tmp.write_bytes(data)

    # Определяем тип по заголовку первой страницы
    file_kind = "unknown"
    try:
        import openpyxl
        wb = openpyxl.load_workbook(tmp, data_only=True, read_only=True)
        ws = wb[wb.sheetnames[0]]
        header_cells = []
        for r in range(1, 6):
            for c in range(1, 6):
                v = ws.cell(r, c).value
                if v: header_cells.append(str(v).lower())
        header = " | ".join(header_cells)
        if "ведомост" in header or "по товарам на склад" in header:
            file_kind = "stocks"
        elif "прайс" in header or "базовый прайс" in header:
            file_kind = "price"
    except Exception as e:
        log.warning(f"xlsx introspect failed: {e}")

    if file_kind == "stocks":
        target = Path("/root/projects/ai-agents-rb/BOT_TG/Ведомость остатков.xlsx")
        kind_label = "ведомость остатков"
    elif file_kind == "price":
        target = Path("/root/projects/ai-agents-rb/BOT_TG/Прайс и остатки.xlsx")
        kind_label = "прайс-лист"
    else:
        # fallback по имени
        if "ведомост" in fname_low:
            target = Path("/root/projects/ai-agents-rb/BOT_TG/Ведомость остатков.xlsx")
            kind_label = "ведомость остатков (по имени файла)"
        else:
            target = Path("/root/projects/ai-agents-rb/BOT_TG/Прайс и остатки.xlsx")
            kind_label = "прайс-лист (по умолчанию)"

    if target.exists():
        shutil.copy2(target, target.with_suffix(".xlsx.bak"))
    target.write_bytes(data)
    await message.answer(
        f"📥 Сохранил как <b>{kind_label}</b>: <code>{target.name}</code> ({len(data)//1024} КБ). Запускаю синк…",
        parse_mode="HTML",
    )

    # Запускаем sync_prices_stocks.py
    venv_py = "/root/projects/ai-agents-rb/bishoprb-agent/.venv/bin/python"
    try:
        result = subprocess.run(
            [venv_py, "sync_prices_stocks.py"],
            cwd="/root/projects/ai-agents-rb/BOT_TG",
            capture_output=True, text=True, timeout=120,
        )
    except subprocess.TimeoutExpired:
        await message.answer("⏰ Скрипт синка таймаутнул (120 сек).")
        return

    out = (result.stdout or "")[-2500:]
    err = (result.stderr or "")[-500:]
    status_emoji = "✅" if result.returncode == 0 else "⚠️"
    msg_lines = [f"{status_emoji} <b>Синк остатков</b> (rc={result.returncode})", "<pre>", out.strip()[-2000:], "</pre>"]
    if err.strip():
        msg_lines.append(f"<i>stderr: {err.strip()[-300:]}</i>")
    await message.answer("\n".join(msg_lines), parse_mode="HTML")

    if result.returncode != 0:
        return

    # Если в выводе сказано «обновлены» хоть что-то ≠ 0 — пушим. Иначе сообщаем.
    import re as _re
    m_pr = _re.search(r"Цены обновлены:\s*(\d+)", out or "")
    n_st = sum(int(x) for x in _re.findall(r"Остатки обновлены:\s*(\d+)", out or ""))
    n_pr = int(m_pr.group(1)) if m_pr else 0
    if n_pr == 0 and n_st == 0:
        await message.answer("ℹ️ Изменений в products.json нет — данные xlsx уже совпадают с TMA. Пушить нечего.")
        return

    # Пушим через shop_tools.shop_publish
    try:
        from services.shop_tools import shop_publish
        publish_result = shop_publish(comment=f"Синк из xlsx: {fname}")
        await message.answer(f"🚀 Опубликовано в TG-BOT.\n<pre>{publish_result[-1500:]}</pre>", parse_mode="HTML")
    except Exception as e:
        log.exception("shop_publish failed")
        await message.answer(f"⚠️ Синк прошёл, но публикация упала: {e}")


@router.message(F.chat.type == "private", Command(commands=["shop", "магазин"]))
async def cmd_shop(message: Message):
    claude_service.reset_shop_history(message.from_user.id)
    is_owner = message.from_user.id == settings.owner_telegram_id
    if is_owner:
        await message.answer(
            "🛍️ Магазин-режим (владелец). Команды:\n"
            "• «найди в магазине Бразилия Серрадо»\n"
            "• «обнови цену 1 кг у Кения АА на 2700»\n"
            "• «обнови описание у Эфиопия Иргачиф: <текст>»\n"
            "• «добавь товар: название, категория, цена 1кг, цена 200г»\n"
            "• «обнови фото у X» (потом пришли фото) ИЛИ пришли фото с подписью\n"
            "• «опубликуй магазин» — пушит в Railway, передеплой за 2 мин\n\n"
            "После прайс-добавления (моносорт/микролот/смесь) предложу добавить в магазин.\n"
            "Выйти из режима — /shop_off."
        )
    else:
        await message.answer(
            "🛍️ Магазин. Что я могу:\n"
            "• «найди в магазине X»\n"
            "• «покажи карточку X»\n"
            "• «какие подкатегории чая»\n\n"
            "Изменения — только у Дмитрия."
        )


@router.message(F.chat.type == "private", Command(commands=["shop_off", "магазин_выкл"]))
async def cmd_shop_off(message: Message):
    claude_service.reset_shop_history(message.from_user.id)
    shop_tools.clear_pending_photo(message.from_user.id)
    await message.answer("Магазин-режим сброшен.")


# ─── Сервисная служба (отдельный бот @rbr_service_bot + Mini App) ───────────

@router.message(F.chat.type == "private", Command(commands=["service", "сервис", "вызов", "заявка"]))
async def cmd_service(message: Message):
    """Раздаёт ссылки на сервисный бот и Mini App.

    @rbr_service_bot — отдельная экосистема под сервисную службу:
    приём заявок, диспетчеризация мастерам, акты с фото, Vision-разбор.
    Здесь Бишоп просто выдаёт правильные ссылки в зависимости от того,
    кто спрашивает.
    """
    from aiogram.types import InlineKeyboardButton, InlineKeyboardMarkup

    text = (
        "🛠 <b>Сервисная служба Roastberry</b>\n\n"
        "Это отдельный бот для <b>приёма и обработки сервисных заявок</b> "
        "по 5 направлениям: Франко, Алеф, Коммерч, Клиенты, HoReCa Machines.\n\n"
        "<b>Что умеет:</b>\n"
        "• Принимает заявки через приложение или чат\n"
        "• Раздаёт мастерам, ставит SLA и эскалирует\n"
        "• Собирает акты с фото\n"
        "• Девид разбирает фото и заполняет акт автоматически\n"
        "• Шлёт ежедневную сводку владельцу\n\n"
        "<b>Открыть в Telegram:</b>"
    )

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(
            text="🚀 Открыть приложение",
            url="https://t.me/rbr_service_bot/service",
        )],
        [InlineKeyboardButton(
            text="🛠 Я техник",
            url="https://t.me/rbr_service_bot?start=technician",
        ),
         InlineKeyboardButton(
            text="📋 Я менеджер",
            url="https://t.me/rbr_service_bot?start=manager",
        )],
        [InlineKeyboardButton(
            text="☎️ Я клиент (подать заявку)",
            url="https://t.me/rbr_service_bot?start=client",
        )],
    ])

    await message.answer(text, reply_markup=kb, parse_mode="HTML")


# ─── Gmail-ассистент ────────────────────────────────────────────────────────

@router.message(F.chat.type == "private", Command(commands=["inbox", "почта", "gmail"]))
async def cmd_inbox(message: Message):
    """Показывает последние письма из Gmail (только метаданные).

    Уровень: 🟢 БЕЗОПАСНЫЙ — только просмотр, без удаления и архивации.

    Аргументы:
        /inbox 20 7  — показать 20 писем за 7 дней (по умолчанию: 15 за 3 дня)
    """
    if message.from_user and message.from_user.id != settings.owner_telegram_id:
        return  # доступ только владельцу

    if not settings.gmail_user or not settings.gmail_app_password:
        await message.answer(
            "📭 Gmail не настроен.\n\n"
            "Чтобы подключить:\n"
            "1. Включи 2FA в Gmail\n"
            "2. Создай App Password: https://myaccount.google.com/apppasswords\n"
            "3. Добавь в .env Бишепа:\n"
            "   <code>GMAIL_USER=твоя_почта@gmail.com</code>\n"
            "   <code>GMAIL_APP_PASSWORD=xxxxxxxxxxxxxxxx</code>\n"
            "4. Перезапусти Бишепа",
            parse_mode="HTML",
        )
        return

    parts = (message.text or "").split()
    try:
        limit = int(parts[1]) if len(parts) > 1 else 15
        since_days = int(parts[2]) if len(parts) > 2 else 3
    except ValueError:
        limit, since_days = 15, 3
    limit = max(1, min(limit, 50))
    since_days = max(1, min(since_days, 30))

    from services.gmail_tools import GmailService, format_inbox_telegram
    gm = GmailService(settings.gmail_user, settings.gmail_app_password)

    status_msg = await message.answer("⏳ Читаю Gmail…")
    try:
        msgs = await gm.list_recent(limit=limit, since_days=since_days)
    except Exception as e:
        log.exception("gmail list_recent failed")
        await status_msg.edit_text(f"❌ Не удалось получить почту: {type(e).__name__}: {str(e)[:200]}")
        return

    # Классифицируем через Claude (whitelisted идут без вызова)
    from services.gmail_classifier import classify_messages, format_inbox_with_categories
    try:
        classes = await classify_messages(msgs)
    except Exception as e:
        log.exception("gmail classify failed")
        # Фолбэк — без категорий
        header = f"📥 Inbox · {settings.gmail_user} · последние {len(msgs)} за {since_days} дн"
        text = format_inbox_telegram(msgs, header=header)
        await status_msg.delete()
        await _send_long(message, text)
        return

    header = f"📥 Inbox · {settings.gmail_user} · {len(msgs)} писем за {since_days} дн"
    text = format_inbox_with_categories(msgs, classes, header=header)
    await status_msg.delete()
    await _send_long(message, text)


@router.message(F.chat.type == "private", Command(commands=["digest", "сводка", "дайджест"]))
async def cmd_digest(message: Message):
    """Сводка по почте за период с группировкой по категориям.

    Аргументы:
        /digest 7  — за 7 дней (по умолчанию: 1 день)
    """
    if message.from_user and message.from_user.id != settings.owner_telegram_id:
        return
    if not settings.gmail_user or not settings.gmail_app_password:
        await message.answer("Gmail не настроен (см. /inbox)")
        return

    parts = (message.text or "").split()
    try:
        since_days = int(parts[1]) if len(parts) > 1 else 1
    except ValueError:
        since_days = 1
    since_days = max(1, min(since_days, 14))

    from services.gmail_tools import GmailService
    from services.gmail_classifier import classify_messages, format_digest
    gm = GmailService(settings.gmail_user, settings.gmail_app_password)

    status_msg = await message.answer(f"📊 Собираю дайджест за {since_days} дн…")
    try:
        msgs = await gm.list_recent(limit=100, since_days=since_days)
    except Exception as e:
        log.exception("digest fetch failed")
        await status_msg.edit_text(f"❌ Не получил почту: {type(e).__name__}: {str(e)[:200]}")
        return

    if not msgs:
        await status_msg.edit_text(f"📊 За последние {since_days} дн писем нет.")
        return

    try:
        classes = await classify_messages(msgs)
    except Exception as e:
        log.exception("digest classify failed")
        await status_msg.edit_text(f"❌ Классификация упала: {type(e).__name__}: {str(e)[:200]}")
        return

    period_label = f"за {since_days} дн" if since_days > 1 else "за сутки"
    text = format_digest(msgs, classes, period_label=period_label)
    await status_msg.delete()
    await _send_long(message, text)


@router.message(F.chat.type == "private", Command(commands=["orders", "заказы", "магазин_сводка"]))
async def cmd_orders_digest(message: Message, bot: Bot):
    """Сводка по заказам магазина за сутки (только владелец)."""
    if message.from_user and message.from_user.id != settings.owner_telegram_id:
        return
    from services.shop_digest import fetch_sync_snapshot, build_digest_text
    snapshot = await fetch_sync_snapshot()
    if snapshot is None:
        await message.answer(
            "⚠️ Не удалось получить данные TG-BOT (/sync). "
            "Проверь TG_BOT_URL и TG_BOT_API_TOKEN в .env Bishop."
        )
        return
    text = build_digest_text(snapshot, since_hours=24)
    await message.answer(text[:4000], parse_mode="HTML")


@router.message(F.chat.type == "private", Command(commands=["health", "checks", "статус"]))
async def cmd_health(message: Message):
    """Сводный статус бизнес-чеков (каталог, Девид, …)."""
    if message.from_user and message.from_user.id != settings.owner_telegram_id:
        return
    from services.business_health import get_status_snapshot
    text = await get_status_snapshot()
    await message.answer(text, parse_mode="HTML", disable_web_page_preview=True)


# ─── Cleanup state (in-memory, per-user) ────────────────────────────────────
# Хранит последний propose: {user_id: {"uids": [...], "categories": {...}, "summary": str}}
_CLEANUP_PENDING: dict[int, dict] = {}


@router.message(F.chat.type == "private", Command(commands=["cleanup", "почистить", "уборка"]))
async def cmd_cleanup(message: Message):
    """Управление чисткой почты. Подкоманды:
        /cleanup            — то же что propose
        /cleanup propose [дней]      — предложить что почистить (без действий)
        /cleanup confirm archive     — архивировать предложенное (убрать из INBOX)
        /cleanup confirm trash       — переместить в корзину (восст. 30 дней)
        /cleanup cancel              — забыть последнее предложение
    """
    if message.from_user and message.from_user.id != settings.owner_telegram_id:
        return
    if not settings.gmail_user or not settings.gmail_app_password:
        await message.answer("Gmail не настроен (см. /inbox)")
        return

    parts = (message.text or "").split()
    sub = parts[1].lower() if len(parts) > 1 else "propose"
    user_id = message.from_user.id

    if sub == "cancel":
        if user_id in _CLEANUP_PENDING:
            del _CLEANUP_PENDING[user_id]
            await message.answer("✅ Предложение отменено.")
        else:
            await message.answer("Нечего отменять.")
        return

    if sub == "confirm":
        action = parts[2].lower() if len(parts) > 2 else None
        if action not in ("archive", "trash"):
            await message.answer(
                "Использование:\n"
                "<code>/cleanup confirm archive</code> — убрать из INBOX (восстановимо)\n"
                "<code>/cleanup confirm trash</code> — в корзину (на 30 дней)",
                parse_mode="HTML",
            )
            return
        pending = _CLEANUP_PENDING.get(user_id)
        if not pending:
            await message.answer(
                "❌ Сначала сделай <code>/cleanup propose</code> — увидишь что Бишеп предлагает почистить.",
                parse_mode="HTML",
            )
            return
        uids = pending["uids"]
        if not uids:
            await message.answer("В предложении ничего не было.")
            return
        from services.gmail_tools import GmailService
        gm = GmailService(settings.gmail_user, settings.gmail_app_password)
        status_msg = await message.answer(f"⏳ Применяю «{action}» к {len(uids)} письмам…")
        try:
            if action == "archive":
                n = await gm.archive(uids)
                what = "Архивировано (убрано из INBOX)"
            else:
                n = await gm.trash(uids)
                what = "В корзину (восстановимо 30 дней)"
        except Exception as e:
            log.exception("cleanup confirm failed")
            await status_msg.edit_text(f"❌ Ошибка: {type(e).__name__}: {str(e)[:200]}")
            return
        del _CLEANUP_PENDING[user_id]
        await status_msg.edit_text(
            f"✅ <b>{what}</b>: {n} писем\n\n"
            f"<i>Восстановить можно из «Архив» или «Корзина» в Gmail.</i>",
            parse_mode="HTML",
        )
        return

    if sub == "propose" or sub.isdigit():
        # /cleanup 14 (без 'propose') — тоже считаем propose
        try:
            since_days = int(parts[2]) if (len(parts) > 2 and sub == "propose") else (
                int(sub) if sub.isdigit() else 7
            )
        except ValueError:
            since_days = 7
        since_days = max(1, min(since_days, 30))

        from services.gmail_tools import GmailService
        from services.gmail_classifier import classify_messages, CATEGORIES

        status_msg = await message.answer(f"⏳ Анализирую почту за {since_days} дней…")
        gm = GmailService(settings.gmail_user, settings.gmail_app_password)
        try:
            msgs = await gm.list_recent(limit=300, since_days=since_days)
        except Exception as e:
            await status_msg.edit_text(f"❌ Ошибка: {type(e).__name__}: {str(e)[:200]}")
            return
        if not msgs:
            await status_msg.edit_text("Писем нет, чистить нечего.")
            return

        classes = await classify_messages(msgs)
        cleanup_categories = ("promo", "spam", "service")
        by_cat: dict[str, list] = {c: [] for c in cleanup_categories}
        cleanup_uids: list[str] = []
        for m, c in zip(msgs, classes):
            if c.category not in by_cat:
                continue
            # Whitelist — никогда не трогаем
            if m.is_whitelisted:
                continue
            by_cat[c.category].append((m, c))
            cleanup_uids.append(m.uid)

        total = sum(len(v) for v in by_cat.values())
        if total == 0:
            await status_msg.edit_text(
                f"✅ Чистить нечего — за {since_days} дн нет промо/спама/уведомлений."
            )
            return

        lines = [
            f"🧹 <b>Можно почистить</b> · за {since_days} дн",
            f"<i>Всего: {total} писем</i> (whitelist-отправители исключены автоматически)",
            "",
        ]
        for cat in cleanup_categories:
            items = by_cat[cat]
            if not items:
                continue
            emo, label = CATEGORIES[cat]
            lines.append(f"<b>{emo} {label}: {len(items)}</b>")
            for m, c in items[:5]:
                sender = (m.from_name or m.from_email)[:30].replace("<", "&lt;").replace(">", "&gt;")
                subj = (m.subject or "(без темы)")[:50].replace("<", "&lt;").replace(">", "&gt;")
                lines.append(f"  • {sender} — {subj}")
            if len(items) > 5:
                lines.append(f"  <i>… ещё {len(items) - 5}</i>")
            lines.append("")

        lines.append(
            "👉 <b>Что дальше:</b>\n"
            "<code>/cleanup confirm archive</code> — убрать из INBOX (остаются в Архиве, восстановимо)\n"
            "<code>/cleanup confirm trash</code> — в корзину (восстановимо 30 дней)\n"
            "<code>/cleanup cancel</code> — забыть это предложение"
        )

        # Сохраняем pending
        _CLEANUP_PENDING[user_id] = {
            "uids": cleanup_uids,
            "since_days": since_days,
            "total": total,
        }

        await status_msg.delete()
        await _send_long(message, "\n".join(lines))
        return

    await message.answer(
        "Использование:\n"
        "<code>/cleanup [дней]</code> — что можно почистить\n"
        "<code>/cleanup confirm archive</code> — архивировать\n"
        "<code>/cleanup confirm trash</code> — в корзину\n"
        "<code>/cleanup cancel</code> — отменить",
        parse_mode="HTML",
    )


# ─── Uptime-мониторинг ──────────────────────────────────────────────────────

@router.message(F.chat.type == "private", Command(commands=["monitor", "monitors", "монитор"]))
async def cmd_monitor(message: Message):
    """Управление uptime-мониторами.

    Подкоманды:
        /monitor                       — список всех
        /monitor add <url> [имя]       — добавить
        /monitor pause <id>            — поставить на паузу
        /monitor resume <id>           — снять с паузы
        /monitor remove <id>           — удалить
    """
    if message.from_user and message.from_user.id != settings.owner_telegram_id:
        return

    from services import uptime_service

    parts = (message.text or "").split(maxsplit=3)
    sub = parts[1].lower() if len(parts) > 1 else ""

    if not sub:
        mons = await uptime_service.list_monitors()
        await message.answer(
            uptime_service.format_monitors_for_telegram(mons),
            parse_mode="HTML",
            disable_web_page_preview=True,
        )
        return

    if sub == "add":
        if len(parts) < 3:
            await message.answer("Использование: <code>/monitor add &lt;url&gt; [имя]</code>", parse_mode="HTML")
            return
        url = parts[2]
        name = parts[3] if len(parts) > 3 else url.replace("https://", "").replace("http://", "").rstrip("/")[:80]
        m = await uptime_service.add_monitor(
            name=name, url=url,
            alert_chat_id=message.chat.id,
            interval_seconds=300,
        )
        await message.answer(
            f"✅ Добавил монитор #{m.id}: <b>{m.name}</b>\n<code>{m.url}</code>\nПервая проверка — в течение минуты.",
            parse_mode="HTML",
            disable_web_page_preview=True,
        )
        return

    if sub in ("pause", "resume"):
        if len(parts) < 3:
            await message.answer(f"Использование: <code>/monitor {sub} &lt;id&gt;</code>", parse_mode="HTML")
            return
        try:
            mid = int(parts[2])
        except ValueError:
            await message.answer("ID должен быть числом")
            return
        ok = await uptime_service.set_active(mid, sub == "resume")
        if ok:
            await message.answer(f"✅ Монитор #{mid} {'возобновлён' if sub == 'resume' else 'на паузе'}")
        else:
            await message.answer(f"❌ Монитор #{mid} не найден")
        return

    if sub == "remove":
        if len(parts) < 3:
            await message.answer("Использование: <code>/monitor remove &lt;id&gt;</code>", parse_mode="HTML")
            return
        try:
            mid = int(parts[2])
        except ValueError:
            await message.answer("ID должен быть числом")
            return
        ok = await uptime_service.remove_monitor(mid)
        await message.answer(f"{'✅ Удалил' if ok else '❌ Не найден'} монитор #{mid}")
        return

    await message.answer(
        "Подкоманды:\n"
        "<code>/monitor</code> — список\n"
        "<code>/monitor add &lt;url&gt; [имя]</code>\n"
        "<code>/monitor pause &lt;id&gt;</code> / <code>/monitor resume &lt;id&gt;</code>\n"
        "<code>/monitor remove &lt;id&gt;</code>",
        parse_mode="HTML",
    )


@router.message(F.chat.type == "private", Command(commands=["gmail_check", "gmail_test"]))
async def cmd_gmail_check(message: Message):
    """Быстрая проверка подключения к Gmail."""
    if message.from_user and message.from_user.id != settings.owner_telegram_id:
        return
    if not settings.gmail_user or not settings.gmail_app_password:
        await message.answer("Gmail не настроен (см. /inbox)")
        return
    from services.gmail_tools import GmailService
    gm = GmailService(settings.gmail_user, settings.gmail_app_password)
    ok, reason = await gm.check_login()
    if ok:
        await message.answer(f"✅ Gmail подключение работает: {settings.gmail_user}")
    else:
        await message.answer(f"❌ Не подключился к Gmail:\n<code>{reason}</code>", parse_mode="HTML")


@router.message(F.chat.type == "private", ~F.text.startswith("/"))
async def handle_private_text(message: Message, bot: Bot):
    """Обработка произвольных сообщений в личке — через Claude понимаем что хочет."""
    if not message.text:
        return

    is_owner = message.from_user.id == settings.owner_telegram_id

    # ─── Двухступенчатый роутинг намерений ───────────────────────────────
    # 1) Быстрый keyword-detector — почти бесплатно
    # 2) Claude Haiku-классификатор как fallback (если keyword не сработал)
    has_shop_history = claude_service.has_active_shop_history(message.from_user.id)
    has_price_history = claude_service.has_active_price_history(message.from_user.id)

    kw_service_admin = is_owner and _looks_like_service_admin_request(message.text)
    kw_service = _looks_like_service_request(message.text)
    kw_shop = _looks_like_shop_request(message.text)
    kw_price = _looks_like_price_request(message.text)
    kw_gmail = is_owner and _looks_like_gmail_request(message.text)

    # Если ни один keyword не сработал и нет активной истории — спросим Claude.
    intent = None
    if not (kw_service or kw_shop or kw_price or kw_gmail or kw_service_admin
            or has_shop_history or has_price_history):
        if 5 <= len(message.text) <= 250:
            try:
                intent = await claude_service.classify_intent(message.text)
            except Exception as e:
                log.warning(f"classify_intent error: {e}")

    # Админ-команды над сервисом (только владелец) идут в shop_chat —
    # там tool-use с service_move_call/issue_code/users_stats/email_route.
    # Имеет приоритет над плашкой SERVICE.
    if kw_service_admin:
        try:
            answer, photos = await claude_service.shop_chat(
                message.text, message.from_user.id, is_owner=is_owner,
            )
        except Exception as e:
            log.exception("shop_chat (service-admin) error: %s", e)
            answer, photos = f"Ошибка: {e}", []
        if photos:
            await _send_shop_photos(message, photos)
        if answer:
            await _send_long(message, answer)
        return

    # Сервис — самый высокий приоритет (не пускает в магазин/прайс ни одно
    # упоминание ремонта).
    if kw_service or intent == "SERVICE":
        await message.answer(
            "🛠 <b>Сервисная служба Roastberry</b>\n\n"
            "Это отдельный бот: приём заявок, мастера, акты с фото, статистика.\n"
            "Открывай через ссылки ниже — они уже зашиты под нужную роль.",
            reply_markup=_service_links_keyboard(),
            parse_mode="HTML",
        )
        return

    # Магазин / курсы / Gmail — все идут в shop_chat (там Claude tool-use
    # сам выбирает что делать).
    gmail_intent = kw_gmail or intent == "GMAIL"
    if (has_shop_history or kw_shop or gmail_intent
            or intent == "SHOP"):
        try:
            answer, photos = await claude_service.shop_chat(
                message.text, message.from_user.id, is_owner=is_owner,
            )
        except Exception as e:
            log.error(f"shop_chat failed: {e}")
            answer, photos = f"Ошибка: {e}", []
        if photos:
            await _send_shop_photos(message, photos)
        if answer:
            await _send_long(message, answer)
        return

    # Прайс-режим. Owner — полный доступ; остальные — только просмотр.
    if has_price_history or kw_price or intent == "PRICE":
        try:
            answer, files = await claude_service.price_chat(
                message.text,
                message.from_user.id,
                is_owner=is_owner,
            )
        except Exception as e:
            log.error(f"price_chat failed: {e}")
            answer, files = f"Ошибка: {e}", []
        if files:
            await _send_price_files(message, files)
        if answer:
            await _send_long(message, answer)
        return

    async with async_session_maker() as session:
        user = await _ensure_user(session, message.from_user)
        if not user.has_started_dm:
            user.has_started_dm = True
            await session.commit()

        # Собираем открытые задачи пользователя
        tasks = await task_service.get_pending_tasks_for_user(
            session, message.from_user.id
        )

        # Явный intent на НОВОЕ напоминание/задачу — пропускаем проверку
        # существующих задач. Идём в shop_chat (там подключён reminder_create
        # tool), а не в general_chat (там tools нет, Claude отказывается).
        text_lower = (message.text or "").lower()
        NEW_REMINDER_KEYWORDS = (
            "напомин", "напомни ", "напомнить",
            "запиши", "запиш",
            "новое", "новую", "новый",
            "в календарь", "в календ",
            "не забыть", "не забуду",
            "сделай напом", "поставь напом",
        )
        is_new_intent = (
            not text_lower.startswith("#")
            and any(kw in text_lower for kw in NEW_REMINDER_KEYWORDS)
        )

        if is_new_intent:
            # Используем shop_chat (там tools для reminder_create + остальное).
            try:
                answer, photos = await claude_service.shop_chat(
                    message.text, message.from_user.id, is_owner=is_owner,
                )
            except Exception as e:
                log.error(f"shop_chat (reminder intent) failed: {e}")
                answer, photos = f"Ошибка: {e}", []
            if photos:
                await _send_shop_photos(message, photos)
            if answer:
                await _send_long(message, answer)
            return

        if not tasks:
            # Свободный диалог через Claude — Бишоп умеет отвечать на общие
            # вопросы про Roastberry, переадресует на нужные команды.
            answer = ""
            try:
                answer = await claude_service.general_chat(
                    message.text,
                    message.from_user.id,
                    is_owner=is_owner,
                )
            except Exception as e:
                log.error(f"general_chat failed: {e}")
            if not answer:
                answer = (
                    "У тебя нет открытых задач. Если нужна помощь — "
                    "напиши /что_ты_знаешь."
                )
            await _send_long(message, answer)
            return

        tasks_info = [
            {
                "id": t.id,
                "description": t.description,
                "deadline": t.deadline.strftime("%Y-%m-%d %H:%M"),
            }
            for t in tasks
        ]

        result = await claude_service.understand_completion_reply(
            message.text, tasks_info
        )
        action = result.get("action")

        if result.get("clarification_needed"):
            list_text = "\n".join(
                f"#{t.id} — {t.description} (до {t.deadline.strftime('%d.%m %H:%M')})"
                for t in tasks
            )
            await message.answer(
                f"У тебя несколько открытых задач. К какой относится сообщение?\n\n"
                f"{list_text}"
            )
            return

        task_id = result.get("task_id")
        if not task_id:
            await message.answer(
                "Не понял к какой задаче относится. Напиши /мои_задачи чтобы увидеть список."
            )
            return

        task = next((t for t in tasks if t.id == task_id), None)
        if not task:
            await message.answer("Не нашёл такую задачу.")
            return

        if action == "complete":
            await task_service.complete_task(session, task.id)
            # Уведомляем постановщика
            try:
                await bot.send_message(
                    task.creator_id,
                    f"✅ Задача закрыта\n\n"
                    f"📋 {task.description}\n"
                    f"👤 Исполнитель: {user.display_name}",
                )
            except Exception as e:
                log.warning(f"Couldn't notify creator: {e}")
            await message.answer(f"✅ Принял, задача закрыта:\n📋 {task.description}")
            return

        if action == "postpone":
            new_deadline_iso = result.get("new_deadline_iso")
            if not new_deadline_iso:
                await message.answer(
                    "Не понял новый дедлайн. Напиши конкретнее, например \"перенеси на пятницу 18:00\"."
                )
                return
            new_deadline = datetime.fromisoformat(new_deadline_iso)
            old_deadline = task.deadline
            await task_service.postpone_task(session, task.id, new_deadline)
            # Уведомляем постановщика
            try:
                await bot.send_message(
                    task.creator_id,
                    f"📅 Перенос дедлайна\n\n"
                    f"📋 {task.description}\n"
                    f"👤 Исполнитель: {user.display_name}\n"
                    f"Было: {old_deadline.strftime('%d.%m %H:%M')}\n"
                    f"Стало: {new_deadline.strftime('%d.%m %H:%M')}\n\n"
                    f"Если хочешь отклонить — напиши мне \"отклонить перенос #{task.id}\".",
                )
            except Exception as e:
                log.warning(f"Couldn't notify creator: {e}")
            await message.answer(
                f"📅 Перенёс дедлайн\n\n"
                f"📋 {task.description}\n"
                f"Новый дедлайн: {new_deadline.strftime('%d.%m %H:%M')}\n\n"
                f"Постановщик уведомлён."
            )
            return

        if action == "question":
            await message.answer(
                "Если есть вопрос по задаче — лучше спроси у постановщика. "
                "Я только напоминаю и принимаю отчёт о выполнении."
            )
            return

        await message.answer(
            "Не понял что ты хочешь. Напиши \"готово\" или \"перенеси на ...\"."
        )
